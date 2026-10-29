#!/usr/bin/env python3
"""
통합교육연수시스템(중앙교육연수원 운영) 이수현황 엑셀 -> static_data/neti_completion.json (기관 단위 집계, 개인정보 없음)

갱신: 새 엑셀을 받으면  python3 neti_extract.py <엑셀경로>  (openpyxl 필요) 후 커밋.
  - Sheet2: 연도별 전체(집합·혼합·정기·상시 × 승인·이수)
  - YYYY년_상세데이터: 연간 기관별, YYYY년_상세데이터(~06.30): 상반기 기관별
기관명이 바뀐 경우(예: 경기도교육연수원 -> 경기도교육청남부연수원) 현재 이름으로 묶는다(RENAME).
"""
import json
import re
import sys
from pathlib import Path

OUT = Path(__file__).parent / "static_data" / "neti_completion.json"
TYPES = ["집합", "혼합", "정기", "상시"]
# 옛 이름 -> 현재 이름(같은 기관의 개칭·통합으로 확인되는 것만)
RENAME = {
    "강원도교육연수원": "강원특별자치도교육청교육연수원",
    "경기도교육연수원": "경기도교육청남부연수원", "경기도교육청교육연수원": "경기도교육청남부연수원",
    "경기도언어교육연수원": "경기도교육청국제교육원", "경기도국제교육원": "경기도교육청국제교육원",
    "경기도율곡교육연수원": "경기도교육청율곡연수원",
    "경기도융합과학교육원": "경기도교육청미래과학교육원",
    "경기도혁신교육연수원": "경기도미래교육연수원",
    "광주광역시교육연수원": "광주광역시교육청교육연수원",
    "부산광역시교육연수원": "부산광역시교육청교육연수원",
    "부산광역시창의융합교육원": "부산광역시교육청창의융합교육원",
    "세종특별자치시교육원": "세종특별자치시교육청교육원",
    "전라남도교육연수원": "전라남도교육청교육연수원",
    "전라북도교육연수원": "전북특별자치도교육청교육연수원", "전라북도교육청교육연수원": "전북특별자치도교육청교육연수원",
    "충청북도단재교육연수원": "충북단재교육연수원",
}
REGION = [("서울", "서울"), ("부산", "부산"), ("대구", "대구"), ("인천", "인천"), ("광주", "광주"), ("대전", "대전"), ("울산", "울산"),
          ("세종", "세종"), ("경기", "경기"), ("강원", "강원"), ("충청북도|충북|단재", "충북"), ("충청남도|충남", "충남"),
          ("전라북도|전북", "전북"), ("전라남도|전남", "전남"), ("경상북도|경북", "경북"), ("경상남도|경남", "경남"), ("제주|탐라", "제주")]


def region_of(name):
    for pat, r in REGION:
        if re.search(pat, name):
            return r
    return "국립·전국"


def num(v):
    try:
        return int(v or 0)
    except (TypeError, ValueError):
        return 0


def row_vals(r, off):
    """승인·이수 쌍 4유형 + 전체"""
    out = {}
    for i, t in enumerate(TYPES + ["전체"]):
        out[t] = [num(r[off + i * 2]), num(r[off + i * 2 + 1])]
    return out


def main(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    totals = {}
    for r in list(wb.worksheets[0].iter_rows(values_only=True))[3:]:
        if r[0] is None:
            continue
        y = str(r[0])
        key = y[:4] + ("H1" if "6.30" in y else "")
        totals[key] = row_vals(r, 1)
    inst = {}  # 기관 -> {기간키: 값}
    for ws in wb.worksheets[1:]:
        m = re.match(r"(\d{4})년_상세데이터(\(~06\.30\))?", ws.title)
        if not m:
            continue
        key = m.group(1) + ("H1" if m.group(2) else "")
        for r in list(ws.iter_rows(values_only=True))[2:]:
            if not r[1]:
                continue
            name = RENAME.get(str(r[1]).strip(), str(r[1]).strip())
            cur = inst.setdefault(name, {}).setdefault(key, {t: [0, 0] for t in TYPES + ["전체"]})
            for t, (a, b) in row_vals(r, 2).items():  # 같은 해 통합된 옛 이름은 합산
                cur[t][0] += a
                cur[t][1] += b
    OUT.write_text(json.dumps({
        "source": "통합교육연수시스템 이수현황(중앙교육연수원)", "as_of": "2026-06-30",
        "note": "상시=상시 운영 원격연수(연중 수강), 정기=기수제 원격연수. 기관명은 현재 이름으로 묶음.",
        "types": TYPES, "totals": totals,
        "institutions": [{"name": n, "region": region_of(n), "years": v} for n, v in sorted(inst.items())],
        "renamed": RENAME,
    }, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"저장: 연도 {len(totals)} · 기관 {len(inst)} -> {OUT}")


if __name__ == "__main__":
    main(sys.argv[1])
