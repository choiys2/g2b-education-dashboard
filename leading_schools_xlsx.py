#!/usr/bin/env python3
"""
AI교육 선도학교(충북·전남) 연락처 엑셀 -> docs/AI선도학교_충북_전남.xlsx

neis_full_export.py 의 선도학교 명단(odcloud)과 나이스 학교기본정보를 그대로 쓰되,
대상 지역만 걸러 영업용 컬럼(컨택상태/담당자/메모)을 붙여 엑셀로 낸다.
환경변수: NEIS_KEY, ODCLOUD_KEY (neis_full_export.py 와 동일)
"""
import sys, time
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from neis_full_export import NEIS_KEY, KIND_MAP, call_neis, fetch_leading_schools

TARGETS = {"M10": "충북", "Q10": "전남"}
ALIASES = {"충북": "충북", "충청북도": "충북", "충북교육청": "충북", "충청북도교육청": "충북",
           "전남": "전남", "전라남도": "전남", "전남교육청": "전남", "전라남도교육청": "전남"}
GRADE_ALIASES = {"초": "초", "초등학교": "초", "중": "중", "중학교": "중", "고": "고", "고등학교": "고"}
KNOWN = {"소속지역", "학교급", "학교명"}


def norm(name):
    return (name or "").replace(" ", "").strip()


def build_lookup():
    lookup = {}
    for code, region in TARGETS.items():
        for grade, kind in KIND_MAP.items():
            page = 1
            while True:
                data = call_neis("schoolInfo", {"KEY": NEIS_KEY, "Type": "json", "pIndex": page, "pSize": 1000,
                                                 "ATPT_OFCDC_SC_CODE": code, "SCHUL_KND_SC_NM": kind})
                if not data or "schoolInfo" not in data:
                    break
                rows = data["schoolInfo"][1]["row"]
                total = data["schoolInfo"][0]["head"][0]["list_total_count"]
                for r in rows:
                    lookup[(region, grade, norm(r.get("SCHUL_NM")))] = r
                if page * 1000 >= total or not rows:
                    break
                page += 1
            time.sleep(0.1)
    return lookup


def class_count(info):
    data = call_neis("classInfo", {"KEY": NEIS_KEY, "Type": "json", "pIndex": 1, "pSize": 1,
                                    "ATPT_OFCDC_SC_CODE": info["ATPT_OFCDC_SC_CODE"],
                                    "SD_SCHUL_CODE": info["SD_SCHUL_CODE"], "AY": "2026"})
    time.sleep(0.08)
    return data["classInfo"][0]["head"][0]["list_total_count"] if data and "classInfo" in data else None


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else "docs/AI선도학교_충북_전남.xlsx"
    leading = fetch_leading_schools()
    print(f"선도학교 전체 {len(leading)}건", file=sys.stderr)
    if not leading:
        sys.exit("선도학교 명단을 받지 못했습니다 (ODCLOUD_KEY 확인)")
    print("소속지역 값 분포:", Counter(r.get("소속지역") for r in leading).most_common(), file=sys.stderr)

    targets = [r for r in leading if ALIASES.get(norm(r.get("소속지역"))) in TARGETS.values()]
    print(f"충북·전남 {len(targets)}건", file=sys.stderr)
    if not targets:
        sys.exit("충북·전남 선도학교가 0건입니다 (소속지역 값 분포 확인)")
    extra_cols = [k for k in targets[0].keys() if k not in KNOWN]

    lookup = build_lookup()
    rows, unmatched = [], 0
    for r in targets:
        region = ALIASES[norm(r.get("소속지역"))]
        grade = GRADE_ALIASES.get(norm(r.get("학교급")), r.get("학교급"))
        info = lookup.get((region, grade, norm(r.get("학교명"))))
        if not info:
            unmatched += 1
        rows.append({
            "지역": region, "학교급": grade, "학교명": r.get("학교명"),
            "설립": (info or {}).get("FOND_SC_NM", ""),
            "전화번호": ((info or {}).get("ORG_TELNO") or "").strip(),
            "홈페이지": ((info or {}).get("HMPG_ADRES") or "").strip(),
            "도로명주소": (((info or {}).get("ORG_RDNMA") or "") + " " + ((info or {}).get("ORG_RDNDA") or "")).strip(),
            "학급수": class_count(info) if info else None,
            "extra": [r.get(k) for k in extra_cols],
            "matched": bool(info),
        })
    order = {"충북": 0, "전남": 1}
    gorder = {"초": 0, "중": 1, "고": 2}
    rows.sort(key=lambda x: (order.get(x["지역"], 9), gorder.get(x["학교급"], 9), x["학교명"] or ""))
    print(f"나이스 매칭 실패 {unmatched}건", file=sys.stderr)

    write_xlsx(out, rows, extra_cols)
    print(f"saved {out}")


def write_xlsx(out, rows, extra_cols):
    wb = Workbook()
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="1F4E78")
    sales_fill = PatternFill("solid", fgColor="C65911")
    miss_fill = PatternFill("solid", fgColor="FFF2CC")

    # 요약
    ws = wb.active
    ws.title = "요약"
    ws["A1"] = "AI교육 선도학교 연락처 — 충북·전남"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"출처: 공공데이터포털 AI교육 선도학교 지정정보(15091298) + 나이스 학교기본정보 · 생성 {time.strftime('%Y-%m-%d')}"
    ws["A2"].font = Font(color="808080", size=9)
    hdr = ["지역", "초", "중", "고", "합계", "연락처 미매칭"]
    for c, h in enumerate(hdr, 1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center"); cell.border = border
    for i, reg in enumerate(["충북", "전남"], 5):
        sub = [r for r in rows if r["지역"] == reg]
        vals = [reg] + [sum(1 for r in sub if r["학교급"] == g) for g in ("초", "중", "고")] \
            + [len(sub), sum(1 for r in sub if not r["matched"])]
        for c, v in enumerate(vals, 1):
            ws.cell(row=i, column=c, value=v).border = border
    tot = ["합계"] + [f"=SUM({get_column_letter(c)}5:{get_column_letter(c)}6)" for c in range(2, 7)]
    for c, v in enumerate(tot, 1):
        cell = ws.cell(row=7, column=c, value=v); cell.font = Font(bold=True); cell.border = border
    ws["A9"] = "※ 노란 행 = 선도학교 명단의 학교명이 나이스와 정확히 일치하지 않아 연락처를 못 붙인 학교(수동 확인 필요)"
    ws["A9"].font = Font(color="808080", size=9)
    for c in range(1, 7):
        ws.column_dimensions[get_column_letter(c)].width = 14

    # 학교 목록
    ws = wb.create_sheet("선도학교_목록")
    base = ["No", "지역", "학교급", "학교명", "설립", "전화번호", "홈페이지", "도로명주소", "학급수"]
    sales = ["컨택상태", "담당자", "메모"]
    headers = base + extra_cols + sales
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = sales_fill if h in sales else head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center"); cell.border = border
    for i, r in enumerate(rows, 2):
        vals = [i - 1, r["지역"], r["학교급"], r["학교명"], r["설립"], r["전화번호"], r["홈페이지"],
                r["도로명주소"], r["학급수"]] + r["extra"] + ["", "", ""]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.border = border
            if not r["matched"]:
                cell.fill = miss_fill
        url = r["홈페이지"]
        if url:
            cell = ws.cell(row=i, column=base.index("홈페이지") + 1)
            cell.hyperlink = url if url.startswith("http") else "http://" + url
            cell.font = Font(color="0563C1", underline="single")
    last = len(rows) + 1
    status_col = get_column_letter(len(headers) - 2)
    dv = DataValidation(type="list", formula1='"미접촉,연락완료,제안발송,협의중,계약,보류"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{status_col}2:{status_col}{max(last, 2)}")
    widths = {"No": 5, "지역": 6, "학교급": 7, "학교명": 20, "설립": 7, "전화번호": 15, "홈페이지": 32,
              "도로명주소": 45, "학급수": 7, "컨택상태": 11, "담당자": 10, "메모": 30}
    for c, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(h, 14)
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"
    wb.save(out)


if __name__ == "__main__":
    main()
