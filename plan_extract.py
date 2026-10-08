#!/usr/bin/env python3
"""
경영계획 양식(xlsx) → 관리 탭용 계획 JSON (로컬 전용 도구, CI 에서 실행하지 않음).

계획 수치는 공개 저장소에 올리지 않는다. 출력 JSON 을 GitHub Secret PLAN_2027_JSON 에 그대로 붙여 넣으면
배포 때 plan2027.py 가 읽어 관리 탭(ADMIN_PASSWORD 별도 암호화) 안에만 싣는다.
- 읽는 시트: 1.매출 · 5-1.광고선전비(상세) · 5-2.지급수수료(상세) · 콘텐츠(실적통) · 8.5개년계획 · 7.인원
- '▷연락망' 등 개인정보 시트는 열지 않는다.

사용: python plan_extract.py <양식.xlsx> [출력.json]
"""
import json
import re
import sys

import openpyxl

SALES_KEY = [(r"B2C|개인", "b2c"), (r"B2S|단체", "b2s"), (r"B2G|위탁", "b2g"), (r"제품매출|교재·단행본", "prod"), (r"상품", "goods"), (r"기타", "etc")]
PL_ROWS = {"rev": r"^매출 합계", "cogs": r"^매출원가 합계", "sga": r"^판매관리비 합계", "labor": r"^1\. 인건비", "dep": r"^2\. 감가상각비",
           "adv": r"^4\. 광고선전비", "fee": r"^5\. 지급수수료", "other": r"^7\. 기타경상비", "op": r"^영업이익 \(K-IFRS", "op1": r"^영업이익 \(1차", "op2": r"^영업이익 \(2차"}
FIVE_ROWS = {"rev": r"^매출 합계", "b2c": r"^1\. 온라인 개인", "b2s": r"^2\. 온라인 단체", "b2g": r"^3\. 용역매출", "etc": r"^4\. 제품·상품·기타",
             "cogs": r"^매출원가 합계", "gp": r"^매출총이익", "sga": r"^판매관리비 합계", "labor": r"^1\. 인건비", "adv": r"^3\. 광고선전비",
             "fee": r"^4\. 지급수수료", "op": r"^영업이익 \(기존"}


def n(v):
    try:
        return round(float(v))
    except (TypeError, ValueError):
        return 0


def detail(ws, item_col_rx):
    """'상세 입력' 행(컴퍼니=콘텐츠컴퍼니) → [{name, basis, prev, plan, m[12]}]"""
    out = []
    for r in ws.iter_rows(values_only=True):
        c = list(r)
        if "콘텐츠컴퍼니" not in [str(x) for x in c[1:3]] or c[2] != "콘텐츠":
            continue
        txt = [str(x) if x is not None else "" for x in c]
        # 마지막 14칸 = 전년 예상, 계획 합계, 1~12월
        tail = c[-14:]
        name = next((t for t in txt[5:10] if t and not re.match(r"^(매출|판관|연간|온라인매출|용역매출|제품매출|상품매출|기타매출|광고선전비|지급수수료)", t)
                     and not t.endswith("Cell") and not re.match(item_col_rx, t)), "")
        basis = next((t for t in txt[6:12] if re.search(r"→|×|%|수준|월 ", t)), "")
        out.append({"name": name[:40], "basis": basis[:80], "prev": n(tail[0]), "plan": n(tail[1]), "m": [n(x) for x in tail[2:]]})
    return out


def summary_rows(ws, pats, ncols):
    out = {}
    for r in ws.iter_rows(values_only=True):
        c = list(r)
        lab = next((str(x).strip() for x in c[:3] if isinstance(x, str) and x.strip()), "")
        for k, rx in pats.items():
            if k not in out and re.search(rx, lab):
                i = [j for j, x in enumerate(c) if isinstance(x, str) and x.strip() == lab][0]
                out[k] = [n(x) for x in c[i + 1:i + 1 + ncols]]
    return out


def main():
    src = sys.argv[1]
    dst = sys.argv[2] if len(sys.argv) > 2 else "plan_2027.local.json"
    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
    sales = []
    for d in detail(wb["1.매출"], r"^교원연수$"):
        k = next((k for rx, k in SALES_KEY if re.search(rx, d["name"])), "etc")
        sales.append({"k": k, **d})
    pl = {k: {"prev": v[0], "plan": v[1], "m": v[2:14]} for k, v in summary_rows(wb["콘텐츠"], PL_ROWS, 14).items()}
    five = summary_rows(wb["8.5개년계획"], FIVE_ROWS, 6)
    heads = 0
    for r in wb["7.인원"].iter_rows(values_only=True):
        if r and any(isinstance(x, str) and x.strip() == "합 계" for x in r[:3]):
            heads = n([x for x in r if isinstance(x, (int, float))][-1])
            break
    r_etc = None
    for r in wb["8.5개년계획"].iter_rows(values_only=True):
        m = next((re.search(r"제품·상품·기타매출의\s*([\d.]+)%", str(x)) for x in r if x and re.search(r"제품·상품·기타매출의\s*[\d.]+%", str(x))), None)
        if m:
            r_etc = float(m.group(1)) / 100
            break
    plan = {"v": 1, "year": 2027, "unit": "원", "heads": heads, "r_etc": r_etc, "sales": sales,
            "ad": detail(wb["5-1.광고선전비(상세)"], r"^광고선전비"), "fees": detail(wb["5-2.지급수수료(상세)"], r"^지급수수료"),
            "pl": pl, "five": {"years": [2026, 2027, 2028, 2029, 2030, 2031], **five}}
    txt = json.dumps(plan, ensure_ascii=False, separators=(",", ":"))
    open(dst, "w", encoding="utf-8").write(txt)
    print(f"계획 JSON {len(txt):,}자 → {dst} (매출 {len(sales)}행 · 광고 {len(plan['ad'])} · 수수료 {len(plan['fees'])} · 인원 {heads})")


if __name__ == "__main__":
    main()
