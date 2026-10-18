r"""교육청 연수 결과보고 자동화 — 원본 1개 + 입력값 1개 → 제출 서류 일체.

사용법
  1) 입력값 양식 만들기(처음 한 번):
       python build_result_report.py --init work/입력값.xlsx
  2) 실행:
       python build_result_report.py \
         --raw work/GG2504-15.xlsx \            # RMS '3.결과보고 자료' 원본
         --input work/입력값.xlsx \             # 사업명·예산·단가·만족도·후기
         --analysis-template work/이수자분석양식.xlsx \
         --invoice-template work/청구서양식.xls \
         --out work/out

산출물(out/)
  결과보고서_<기수>.hwpx        공문 + 붙임1 결과보고서
  붙임2_이수자명단_<기수>.xlsx   이수자만, 휴대폰번호 제외(제출용)
  붙임3_청구서_<기수>.xls/.pdf   학점별 단가 × 이수 인원, 예산 초과분 교육기부
  이수자분석_<기수>.xlsx         분석양식에 정리본 붙여넣기(엑셀로 열면 피벗·그래프 자동 갱신)
  0_데이터점검_<기수>.xlsx       정리본·변경내역·점검목록·문자요청 대상·경영분석
"""

import argparse
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import rr_clean  # noqa: E402
import rr_hwpx  # noqa: E402
import rr_input  # noqa: E402
import rr_stats  # noqa: E402
import rr_xlsx_template  # noqa: E402

LIST_COLUMNS = [c for c in rr_clean.COLUMNS if c != "휴대폰번호"]
LIST_WIDTHS = [10.6, 9, 16.3, 51.5, 12.9, 25.1, 7, 10, 14.5, 24.4, 21.1, 30.9, 10.1, 9, 14.8]
THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = dict(font=Font(name="맑은 고딕", size=11, bold=True), fill=PatternFill("solid", fgColor="00B0F0"),
            alignment=Alignment(horizontal="center", vertical="center"), border=BOX)
BODY_FONT = Font(name="맑은 고딕", size=11)


def _sheet(ws, header, rows, widths=None):
    ws.append(header)
    for c in ws[1]:
        for k, v in HEAD.items():
            setattr(c, k, v)
    ws.row_dimensions[1].height = 29.25
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.border = BODY_FONT, BOX
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    if rows:
        ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(header))}{len(rows) + 1}"
    ws.freeze_panes = "A2"


def write_completer_list(path, rows, sheet_name):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (sheet_name or "이수자 명단")[:31]
    done = [[r[h] for h in LIST_COLUMNS] for r in rows if r["점수"] == "이수"]
    _sheet(ws, LIST_COLUMNS, done, LIST_WIDTHS)
    ws.freeze_panes = None
    wb.save(path)
    return len(done)


def write_check_book(path, rows, log, issues, s, inv, basic, cohort):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "정리본"
    _sheet(ws, rr_clean.COLUMNS, [[r[h] for h in rr_clean.COLUMNS] for r in rows],
           [10.6, 9, 10, 45, 12.9, 22, 8, 10, 13, 14, 26, 20, 18, 10, 8, 10])

    sms = [r for kind, r, _ in issues if kind == "필수·문자요청"]
    ws = wb.create_sheet("문자요청_나이스번호")
    msg = ("[비상 티스쿨] {name} 선생님, '{course}' 연수 이수를 축하드립니다. 이수 등록을 위해 "
           "나이스 개인번호(J로 시작하는 10자리)를 회신 부탁드립니다. 문의 1544-9044")
    _sheet(ws, ["성명", "휴대폰번호", "학교명", "과정명", "이수번호", "발송문구(VENS 대량전송용)"],
           [[r["성명"], r["휴대폰번호"], r["학교명"], r["과정명"].strip(), r["이수번호"],
             msg.format(name=r["성명"], course=r["과정명"].strip())] for r in sms],
           [8, 15, 18, 45, 30, 90])

    ws = wb.create_sheet("점검목록")
    _sheet(ws, ["구분", "원본행", "성명", "과정명", "점수", "내용"],
           [[k, r["_원본행"], r["성명"], r["과정명"].strip(), r["점수"], why] for k, r, why in issues],
           [14, 8, 8, 45, 8, 70])

    ws = wb.create_sheet("변경내역")
    _sheet(ws, ["원본행", "성명", "항목", "원본값", "정리값"], [list(x) for x in log], [8, 8, 14, 22, 22])

    ws = wb.create_sheet("과정별")
    _sheet(ws, ["과정명", "학점", "참여", "이수", "미이수", "이수율"],
           [[c["과정명"].strip(), c["학점"], c["참여"], c["이수"], c["미이수"], round(c["이수율"], 4)] for c in s["과정"]],
           [50, 6, 8, 8, 8, 8])
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
        row[0].number_format = "0.0%"

    ws = wb.create_sheet("경영분석")
    for line in insights(s, inv, issues, log, basic, cohort):
        ws.append([line])
    ws.column_dimensions["A"].width = 140
    wb.save(path)
    return len(sms)


def insights(s, inv, issues, log, basic, cohort):
    """교육청 제출용이 아닌 내부 판단용(재계약·단가·운영 개선)."""
    tot, ppl = s["총"], s["인원"]
    price1 = inv["내역"][0]["단가"] if inv["내역"] else 0
    low = [c for c in s["과정"] if c["참여"] >= 10 and c["이수율"] < tot["이수율"] - 0.05]
    blank_pos = sum(1 for x in log if x[2] == "직위" and x[3] == "")
    lines = [
        f"[{cohort}] 내부 분석 — 교육청 제출용 아님",
        "",
        "① 핵심요약",
        f"  - 신청 {tot['참여']:,}건(실인원 {ppl['신청인원']:,}명) → 이수 {tot['이수']:,}건, 이수율 {tot['이수율']:.1%}",
        f"  - 청구 {inv['청구금액']:,}원 / 산정 {inv['산정금액']:,}원 / 계약 {inv['예산']:,}원 → 예산 소진율 "
        + (f"{inv['청구금액'] / inv['예산']:.1%}" if inv["예산"] else "-"),
    ]
    if inv["교육기부"]:
        lines.append(f"  - 교육기부(초과 이수분 무상) {-inv['교육기부']:,}원 = 이수 {(-inv['교육기부']) // max(price1, 1):,}건분. "
                     "다음 계약 시 예산 증액 근거로 활용")
    elif inv["예산"] and price1:
        lines.append(f"  - 예산 대비 미달 {inv['예산'] - inv['청구금액']:,}원 = 이수 {(inv['예산'] - inv['청구금액']) // price1:,}건 부족."
                     " 종료 전 이수독려 문자 강화 필요")
    lines += [
        "",
        "② 분석(수치·패턴)",
        f"  - 1인당 {ppl['1인당과정수']:.2f}개 과정 수강, 과정 수 분포 {ppl['과정수분포']} → 다과정 패키지 수요 확인",
        f"  - 상위 1개 과정 '{s['과정'][0]['과정명'].strip()}' 비중 {s['과정'][0]['참여'] / tot['참여']:.1%}",
        "  - 평균 이수율보다 5%p 이상 낮은 과정(참여 10건↑): "
        + (", ".join(f"{c['과정명'].strip()[:18]}({c['이수율']:.0%})" for c in low) or "없음"),
        f"  - 미이수 {tot['미이수']:,}건 = 잠재 매출 {tot['미이수'] * price1:,}원(1학점 단가 기준). "
        "진도율 90%↑ 미이수자 별도 독려 시 구제 인원 감소(인계서)",
        f"  - 데이터 품질: 직위 공란 {blank_pos}건, 이수자 나이스번호 누락 "
        f"{sum(1 for k, *_ in issues if k == '필수·문자요청')}건 → 신청 단계 필수입력 강화 필요",
        "",
        "③ 다음 사업 제안 포인트",
        "  - 이수율 높은 과정·참여 상위 과정을 묶은 학교급 맞춤 패키지 제안",
        "  - 초과 이수(교육기부) 발생 시: 실적 근거로 차년도 예산 증액·2학점 과정 확대 제안",
        "  - 이수율 저조 과정: 종료 7일 전·1~3일 전 독려 + 2~3일 간격 추가 발송(인계서 기준)",
    ]
    return lines


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--init", metavar="입력값.xlsx", help="빈 입력값 양식 만들기")
    ap.add_argument("--raw", help="RMS 3.결과보고 자료 원본(xlsx)")
    ap.add_argument("--input", help="입력값.xlsx")
    ap.add_argument("--analysis-template", help="이수자분석양식(xlsx)")
    ap.add_argument("--invoice-template", help="청구서 양식(xls)")
    ap.add_argument("--out", default="out")
    ap.add_argument("--no-pdf", action="store_true", help="청구서 PDF 생략")
    a = ap.parse_args()

    if a.init:
        rr_input.create_blank(a.init)
        print(f"입력값 양식 생성: {a.init}")
        return
    if not (a.raw and a.input):
        ap.error("--raw 와 --input 은 필수입니다 (처음이면 --init 으로 입력값 양식 생성)")

    cfg = rr_input.read(a.input)
    basic = cfg["basic"]
    raw = rr_clean.load_raw(a.raw)
    rows, log = rr_clean.clean(raw, (basic.get("관할교육청(정리 통일값)") or "").strip())
    issues = rr_clean.checks(rows)
    s = rr_stats.summarize(rows)
    cohort = "_".join(s["기수"])

    prices = {1: basic.get("1학점 단가"), 2: basic.get("2학점 단가")}
    need = {c["학점"] for c in s["과정"] if c["이수"]}
    missing = [f"{k}학점 단가" for k in need if not prices.get(k)]
    if missing:
        sys.exit(f"입력값 누락: {', '.join(missing)} (기본정보 시트)")
    inv = rr_stats.invoice(s["과정"], prices, basic.get("계약예산(초기견적가)"))
    sv = rr_stats.survey(cfg["survey_counts"], cfg["respondents"])
    reviews = cfg["reviews_picked"] or rr_input.pick_reviews(cfg["reviews_raw"], basic.get("연수후기 개수") or 8)

    os.makedirs(a.out, exist_ok=True)
    out = lambda name: os.path.join(a.out, name)
    made = []

    n_sms = write_check_book(out(f"0_데이터점검_{cohort}.xlsx"), rows, log, issues, s, inv, basic, cohort)
    made.append(f"0_데이터점검_{cohort}.xlsx")
    n_done = write_completer_list(out(f"붙임2_이수자명단_{cohort}.xlsx"), rows, basic.get("이수자명단 시트명"))
    made.append(f"붙임2_이수자명단_{cohort}.xlsx")

    if a.analysis_template:
        sat = sv["종합"]
        sentence = (f"▣ 학습자는 본 연수를 내용 전문성과 현장 적용성이 높아 추천할 만한 강의로 인식했으며 "
                    f"연수 운영 전체적인 만족도를 보면 {sat['만족비율'] * 100:.1f}%가 '만족'한다고 답했음. "
                    f"5점 만점 기준으로 {sat['점수']:.2f}점의 높은 연수 만족도를 보임")
        rr_xlsx_template.fill(a.analysis_template, out(f"이수자분석_{cohort}.xlsx"), rows, rr_clean.COLUMNS,
                              survey=sv if cfg["respondents"] else None, summary_sentence=sentence)
        made.append(f"이수자분석_{cohort}.xlsx")

    if a.invoice_template:
        import rr_invoice
        with rr_invoice.office() as desk:
            rr_invoice.fill(desk, a.invoice_template, out(f"붙임3_청구서_{cohort}.xls"),
                            None if a.no_pdf else out(f"붙임3_청구서_{cohort}.pdf"), basic, inv)
        made.append(f"붙임3_청구서_{cohort}.xls" + ("" if a.no_pdf else " (+pdf)"))

    rr_hwpx.build(out(f"결과보고서_{cohort}.hwpx"), basic, s, sv, inv, reviews)
    made.append(f"결과보고서_{cohort}.hwpx")

    tot = s["총"]
    print(f"[{cohort}] 참여 {tot['참여']:,}건 · 이수 {tot['이수']:,}건 · 이수율 {tot['이수율']:.2%}")
    print(f"  청구 {inv['청구금액']:,}원 (산정 {inv['산정금액']:,} / 교육기부 {inv['교육기부']:,})")
    print(f"  만족도 {sv['종합']['점수']:.2f}점 ({sv['종합']['백분율']:.1%}) · 후기 {len(reviews)}건")
    print(f"  이수자명단 {n_done:,}건 · 정리 변경 {len(log):,}건 · 점검 {len(issues):,}건 · "
          f"나이스번호 문자요청 {n_sms}명")
    for m in made:
        print(f"  → {os.path.join(a.out, m)}")


if __name__ == "__main__":
    main()
