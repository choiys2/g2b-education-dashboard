"""교육청 결과보고 — 1단계: 시스템 원본(3.결과보고 자료) 정리.

원본은 RMS '성적산출관리 > 이수결과관리 > 3.결과보고 자료 다운' 엑셀(16열).
학습자가 비워 둔 칸·표기 흔들림을 업무 인계서 규칙대로 정리하고,
나이스번호가 없는 이수자(보고 후 문자 재요청 대상)를 따로 뽑는다.

정리 규칙은 GG2504-15(경기 용인 중등) 실제 결과보고본과 원본을 대조해 역산했다.
  - 관할교육청: 사업 관할청 하나로 통일(공란·'경기도교육청'·뒤 공백 포함)
  - 직위: 공란 → 교사, *교사·부장류 → 교사, 교육연구사·교육전문직 → 교육공무직
          (분석양식 직위표 7개 칸: 장학사·교사·교감·교육공무직·교장·특수교사·기타)
  - 학교구분: 공란 → 소속이 교육청이면 '기타', 학교면 '공립'
  - 정렬: 이수 먼저(원본 순서 유지) → 미이수
"""

import re
from collections import Counter

import openpyxl

COLUMNS = ["기수", "시도교육청", "구분", "과정명", "연수시간", "연수기간", "성명", "생년월일",
           "나이스식별번호", "휴대폰번호", "이수번호", "관할교육청", "학교명", "직위", "점수", "학교구분"]

POSITION_SLOTS = ["장학사", "교사", "교감", "교육공무직", "교장", "특수교사", "기타"]
SCHOOL_TYPES = ["공립", "사립", "국립", "기타"]
NEIS_RE = re.compile(r"^[A-Z]\d{9}$")
PHONE_RE = re.compile(r"^01[016789]-?\d{3,4}-?\d{4}$")


def _s(v):
    return "" if v is None else str(v).strip()


def normalize_position(raw, school_level):
    p = _s(raw).replace(" ", "")
    if not p:
        return "교사", "공란→교사"
    if p in POSITION_SLOTS:
        return p, None
    if p.startswith("장학"):
        return "장학사", f"{p}→장학사"
    if p.startswith("교장") or p == "원장":
        return "교장", f"{p}→교장"
    if p.startswith("교감") or p == "원감":
        return "교감", f"{p}→교감"
    if "특수" in p:
        return "특수교사", f"{p}→특수교사"
    if p in ("교육연구사", "교육연구관", "교육전문직", "연구사", "연구관", "주무관", "실무사", "교육공무직원"):
        return "교육공무직", f"{p}→교육공무직"
    if p.endswith("교사") or p.endswith("부장") or p in ("교수", "강사"):
        return "교사", f"{p}→교사"
    return "기타", f"{p}→기타"


def load_raw(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    header = [_s(c.value) for c in ws[1]]
    missing = [c for c in COLUMNS if c not in header]
    if missing:
        raise ValueError(f"원본에 필수 열이 없습니다: {missing} (시스템 '3.결과보고 자료' 파일인지 확인)")
    idx = {h: header.index(h) for h in COLUMNS}
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not any(v not in (None, "") for v in r):
            continue
        rows.append({h: r[idx[h]] for h in COLUMNS})
    return rows


def clean(rows, office):
    """rows: load_raw 결과. office: 통일할 관할교육청명. 반환 (정리행, 변경로그)."""
    out, log = [], []
    for i, src in enumerate(rows, start=2):
        r = {h: _s(src[h]) for h in COLUMNS}
        r["과정명"] = "" if src["과정명"] is None else str(src["과정명"])  # 시스템 과정명 그대로(끝 공백 포함)
        r["_원본행"] = i

        if office and r["관할교육청"] != office:
            log.append((i, r["성명"], "관할교육청", r["관할교육청"], office))
            r["관할교육청"] = office

        pos, why = normalize_position(r["직위"], r["구분"])
        if why:
            log.append((i, r["성명"], "직위", r["직위"], pos))
            r["직위"] = pos

        if not r["학교구분"]:
            new = "기타" if r["구분"] == "교육청" else "공립"
            log.append((i, r["성명"], "학교구분", "", new))
            r["학교구분"] = new
        elif r["학교구분"] not in SCHOOL_TYPES:
            log.append((i, r["성명"], "학교구분", r["학교구분"], "기타"))
            r["학교구분"] = "기타"
        out.append(r)

    out.sort(key=lambda r: r["점수"] != "이수")
    return out, log


def checks(rows):
    """보고서 제출 전 점검 목록. 나이스번호는 이수자 필수정보."""
    issues = []
    seen_cert = Counter(r["이수번호"] for r in rows if r["이수번호"])
    for r in rows:
        done = r["점수"] == "이수"
        if done and not r["나이스식별번호"]:
            issues.append(("필수·문자요청", r, "이수자 나이스식별번호 누락 — 보고 후에도 문자로 받아 추가 제출"))
        elif r["나이스식별번호"] and not NEIS_RE.match(r["나이스식별번호"]):
            issues.append(("확인", r, f"나이스번호 형식 이상({r['나이스식별번호']})"))
        if done and not r["이수번호"]:
            issues.append(("필수", r, "이수자인데 이수번호 없음 — RMS 성적관리 재확인"))
        if r["이수번호"] and seen_cert[r["이수번호"]] > 1:
            issues.append(("확인", r, f"이수번호 중복({r['이수번호']})"))
        if not re.match(r"^(19|20)\d{6}$", r["생년월일"]):
            issues.append(("확인", r, f"생년월일 형식 이상({r['생년월일']})"))
        if done and not PHONE_RE.match(r["휴대폰번호"]):
            issues.append(("확인", r, f"휴대폰번호 형식 이상({r['휴대폰번호']})"))
        if not r["학교명"]:
            issues.append(("확인", r, "학교명 공란"))
    return issues


def credit_of(hours_text):
    """'15시간/1학점' → 1, '30시간/2학점' → 2."""
    m = re.search(r"(\d+)\s*학점", hours_text or "")
    if m:
        return int(m.group(1))
    m = re.search(r"(\d+)\s*시간", hours_text or "")
    return 2 if m and int(m.group(1)) >= 30 else 1
