"""교육청 결과보고 — 담당자 입력값 엑셀(입력값.xlsx) 생성·읽기.

원본 자료 외에 사람이 넣어야 하는 값은 이 파일 한 곳에만 모은다.
  - 기본정보: 사업명·수신처·계약예산(영업 최종 견적가)·학점별 단가·청구일 등
  - 만족도  : RMS 평가관리 > 강의평가 > 통계의 문항별 응답 수
  - 연수후기: 통계 8번(자유 의견)을 통째로 붙여넣기. B열에 O 표시한 것만 쓰거나, 비우면 자동 선별
"""

import datetime as dt
import re

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from rr_stats import SURVEY_ITEMS

BASIC_FIELDS = [
    ("사업명", "", "청구서·공문 제목에 그대로 들어감"),
    ("수신처", "", "예) 경기도 용인교육지원청"),
    ("수신 담당자", "", "예) 홍길동 장학사님"),
    ("참조", "", "공문 참조란(없으면 비움)"),
    ("관할교육청(정리 통일값)", "", "원본 '관할교육청' 열을 이 값으로 통일. 비우면 원본 유지"),
    ("연수대상", "", "예) 용인교육지원청 관내 중·고등학교 교원"),
    ("연수종별", "직무연수", ""),
    ("계약예산(초기견적가)", "", "영업 최종 예산. 산정액이 넘으면 초과분을 '교육기부'로 차감"),
    ("1학점 단가", 30000, "15시간 과정 1인 단가(원)"),
    ("2학점 단가", "", "30시간 과정 1인 단가(원). 2학점 과정이 있으면 필수"),
    ("청구일자", "", "YYYY-MM-DD"),
    ("공급자 담당자", "", "청구서 공급자 담당자"),
    ("공급자 이메일", "", ""),
    ("공급자 전화", "02-6970-6204", ""),
    ("공급자 주소", "경기도 과천시 과천대로 2길 54 그라운드브이", ""),
    ("대표자명", "양태회", ""),
    ("등록번호", "211-87-07735", ""),
    ("이수자명단 시트명", "", "예) 경기 용인 중등 이수자. 비우면 '이수자 명단'"),
    ("연수후기 개수", 8, "보고서에 넣을 후기 수(자동 선별 시)"),
]


def create_blank(path):
    wb = openpyxl.Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F4E79")

    ws = wb.active
    ws.title = "기본정보"
    ws.append(["항목", "값", "설명"])
    for k, v, d in BASIC_FIELDS:
        ws.append([k, v, d])
    for c in ws[1]:
        c.font, c.fill = head, fill
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 60

    ws = wb.create_sheet("만족도")
    ws.append(["조사인원", None, "← 통계 화면의 응답자 수"])
    ws.append([])
    ws.append(["구분", "문항", "매우 그렇다", "그렇다", "보통", "그렇지 않다", "매우 그렇지 않다"])
    for name, q in SURVEY_ITEMS:
        ws.append([name, q, None, None, None, None, None])
    ws.append([])
    ws.append(["※ 추천 지수는 보기 3개 설문이라 1·2·3번 응답 수를 각각 '매우 그렇다·그렇다·보통' 칸에 입력(인계서 규칙)"])
    for c in ws[3]:
        c.font, c.fill = head, fill
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 52

    ws = wb.create_sheet("연수후기")
    ws.append(["후기 원문(통계 8번 응답을 한 줄에 하나씩 붙여넣기)", "선택(O)"])
    for c in ws[1]:
        c.font, c.fill = head, fill
    ws.column_dimensions["A"].width = 100
    ws["A2"].alignment = Alignment(wrap_text=True)
    wb.save(path)


def _num(v, field):
    if v in (None, ""):
        return None
    try:
        return int(float(str(v).replace(",", "").replace("원", "").strip()))
    except ValueError:
        raise ValueError(f"입력값 '{field}'이(가) 숫자가 아닙니다: {v!r}")


def read(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    basic = {}
    for row in wb["기본정보"].iter_rows(min_row=2, values_only=True):
        if row and row[0]:
            basic[str(row[0]).strip()] = row[1]
    for f in ("계약예산(초기견적가)", "1학점 단가", "2학점 단가", "연수후기 개수"):
        basic[f] = _num(basic.get(f), f)
    d = basic.get("청구일자")
    if isinstance(d, dt.datetime):
        d = d.date()
    elif isinstance(d, str) and d.strip():
        d = dt.date.fromisoformat(d.strip().replace(".", "-").replace("/", "-"))
    basic["청구일자"] = d or dt.date.today()

    ws = wb["만족도"]
    respondents = _num(ws["B1"].value, "조사인원")
    counts = {}
    for row in ws.iter_rows(min_row=4, max_row=3 + len(SURVEY_ITEMS), values_only=True):
        if row[0]:
            counts[str(row[0]).strip()] = [_num(v, f"만족도 {row[0]}") or 0 for v in row[2:7]]

    lines, picked = [], []
    for a, b in wb["연수후기"].iter_rows(min_row=2, max_col=2, values_only=True):
        if a in (None, ""):
            continue
        for line in str(a).splitlines():
            line = line.strip()
            if line:
                lines.append(line)
                if b and str(b).strip().upper() in ("O", "0", "Y", "V", "○"):
                    picked.append(line)
    return {"basic": basic, "survey_counts": counts, "respondents": respondents,
            "reviews_raw": lines, "reviews_picked": picked}


_JUNK = re.compile(r"^(없|x$|n+$|감사|수고|고생|좋|굿|\W+$|\.+$|-+$)", re.I)
_GOOD = ("현장", "수업", "활용", "적용", "실용", "도움", "유익", "학생", "바로", "쉽", "구체", "사례", "에듀테크", "AI")


def pick_reviews(lines, n):
    """통계 8번 자유 의견에서 '연수의 장점이 드러나는' 후기를 고른다.
    단답·형식적 인사('감사합니다', '없음', '.')는 제외하고 현장 적용·활용 언급을 우선."""
    seen, cands = set(), []
    for t in lines:
        key = re.sub(r"\s+", "", t)
        if key in seen or len(key) < 12:
            continue
        seen.add(key)
        if _JUNK.match(t) and len(key) < 20:
            continue
        if re.search(r"(불편|어렵|아쉽|오류|안\s?됨|안돼|느리|개선|힘들)", t):
            continue
        score = sum(2 for w in _GOOD if w in t) + min(len(key), 80) / 20
        cands.append((score, t))
    cands.sort(key=lambda x: -x[0])
    return [t for _, t in cands[:n]]
